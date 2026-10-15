from pathlib import Path
import json
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent
data=json.loads((OUT/'业务流程信息梳理.json').read_text(encoding='utf-8'))
steps={s['number']:s for s in data['steps']}
records={r['id']:r for r in json.loads((ROOT/'执行结果.json').read_text(encoding='utf-8-sig'))}

# This structure follows the user's illustrated navigation hierarchy.
structure=[
 ('登录系统', [('登录系统',[1])]),
 ('仪表盘', [('查看首页概览',[101])]),
 ('教学计划管理',[
   ('计划准备',[5,6,22,93,94,95,96]),
   ('课程录入与导入',[23,24,25,27,28,30,31,32,34,35,26]),
   ('计划工作量核对',[29,33]),
   ('本科计划提交与审批',[21,36,37,38,39,40,41,42,43,44,45]),
   ('研究生计划提交与审批',[46,47,48,49,50,51,52,53,54]),
 ]),
 ('结算管理',[
   ('SIS课程与课时导入',[55,56,57,58,59]),
   ('结算生成与计算',[60,61,62,63,64,65,66]),
   ('工作量调整与汇总',[67,68,69]),
   ('合教配置与权重',[70,71,72,73,74]),
   ('特殊审批',[89,90,91,92]),
   ('报表导出',[75,76,77,78]),
   ('快照与归档',[79,80,81,82]),
 ]),
 ('统计与查询',[
   ('达标情况对比',[83,84,85]),
   ('工作量查询',[86,87,88]),
 ]),
 ('待办与消息',[
   ('待办任务',[97]),
   ('系统消息',[98,99,100]),
 ]),
 ('系统配置管理',[
   ('学年学期管理',[2,3,4,110]),
   ('计算规则管理',[7,8,9,10,11,12]),
   ('课程信息维护',[13,14,15,16]),
   ('教师资料维护',[17,18,19,20]),
 ]),
 ('验收数据收尾',[
   ('清理模拟开课资料',[102,103,108]),
   ('清理模拟结算资料',[104,105]),
   ('清理模拟基础资料',[106,107,109]),
 ]),
]
order=[i for _,subs in structure for _,ids in subs for i in ids]
assert sorted(order)==list(range(1,111)) and len(set(order))==110

expected={
1:'账号登录成功，进入首页并显示可访问的业务菜单。',
2:'测试学年创建成功，原当前学年保持不变。',
3:'新学年生成Term1、Term2和Summer三个学期。',
4:'学期名称及日期保存成功，重新打开仍一致。',
5:'新学年显示上一学年复制的各合同类型标准。',
6:'测试学年的标准修改成功，不影响原业务学年。',
7:'规则成功保存，名称、条件和系数可回读。',
8:'匹配条件返回配置系数；不匹配条件不命中。',
9:'修改后的系数保存成功，编辑及列表显示一致。',
10:'预览仅展示当前版本中符合规则条件的课程。',
11:'自建规则删除后，不再出现在规则组中。',
12:'原版本规则保持原配置，说明文字与实际权重一致。',
13:'课程新增成功，可以按唯一代码检索。',
14:'名称、课程标记及学分保存后正确回显。',
15:'文件中的课程导入成功，关键字段与文件一致。',
16:'相同代码更新原课程，不重复新增；允许格式有明确说明。',
17:'教师创建成功，工号、姓名、学部和合同类型可查询。',
18:'教师姓名修改成功，重新查询显示新姓名。',
19:'文件中的模拟教师导入成功，资料与文件一致。',
20:'教师身份移除成功，列表不再返回该教师。',
21:'本科助理登录成功，可以进入本科开课页面。',
22:'已配置的具体学期可直接选择，无需先访问其他页面。',
23:'必填项为空时拦截保存，并提示需填写的字段。',
24:'课程、教师、学分及上课安排保存成功，可重新查看。',
25:'标题修改成功，列表显示修改后的内容。',
26:'课程保存后生成可追溯的版本记录。',
27:'复制生成新记录，并保留原课程资料。',
28:'仅删除选中的复制项，保留原模拟课程。',
29:'计划汇总与已保存安排一致，列表课时正确显示。',
30:'标准模板可下载、打开并用于填写导入资料。',
31:'标准模板解析成功，正确显示有效行及错误提示。',
32:'有效课程导入成功，教师、权重及课时可回读。',
33:'每周与隔周正确折算，教师分配与总量一致。',
34:'目标学期生成对应课程，保留学科部等关键字段。',
35:'历史文件导入后显示本次文件、时间和处理结果。',
36:'提交成功，进入待审批状态，助理不可继续改动。',
37:'主任待办显示本轮本科计划和正确课程数。',
38:'原因未填写时，不允许提交驳回。',
39:'计划被驳回，保存处理人、时间及原因。',
40:'助理可重新维护课程并提交，计划重新进入审批。',
41:'主任通过后，流程进入对应学院行政审批。',
42:'本科行政仅处理对应范围，跳转保持计划筛选条件。',
43:'本科行政通过后，流程进入学院行政主管审批。',
44:'行政主管通过后，流程进入主管领导终审。',
45:'主管领导通过后，本科计划最终显示已通过。',
46:'研究生助理可选择具体学期，列表按层次隔离。',
47:'研究生开课资料保存成功，不混入本科计划。',
48:'研究生计划提交成功，助理页面进入只读。',
49:'主任通过后，进入研究生学院行政审批。',
50:'研究生行政可查看本轮对应计划，处理范围正确。',
51:'研究生行政通过后，进入学院行政主管审批。',
52:'行政主管通过后，进入主管领导终审。',
53:'主管领导通过后，研究生计划最终显示已通过。',
54:'两条最终状态正确，流程图完整展示实际审批节点。',
55:'识别缺少关键列的文件，并阻止无效数据导入。',
56:'有效SIS数据按对应学期导入并可查询。',
57:'标准学分列正确解析，相同Class更新且不重复新增。',
58:'错误格式给出提示，无效行不能确认导入。',
59:'按Class编号匹配并更新上课时间与课时。',
60:'按教师拆分记录，完整承接计划课时及工作量。',
61:'缺时间记录明确提示并跳过，不作为正常结果。',
62:'零学分记录明确标为未通过，不生成有效工作量。',
63:'有效样本计算成功，原始课时和最终工作量正确。',
64:'结算结果和偏差说明保存成功，刷新后保留。',
65:'重新计算及保存后，已有偏差说明仍保留。',
66:'时间线只显示该记录真实发生的处理事件。',
67:'调整结果和原因保存，原始课时保持不变。',
68:'教师汇总与已保存的课程数、工作量合计一致。',
69:'明细与保存结果一致，并正确标明异常状态。',
70:'模拟合教代码新增成功并可查询。',
71:'模拟合教代码更新成功并可回读。',
72:'模拟合教代码删除后不再显示。',
73:'教师占比超过100%时，不允许保存。',
74:'合法比例可保存；前置数据不满足时明确提示原因。',
75:'教师明细文件成功生成、可打开且内容正确。',
76:'学部汇总文件成功生成、可打开且内容正确。',
77:'差异汇总文件成功生成、可打开且内容正确。',
78:'无数据时明确提示，不生成误导性的报表。',
79:'基于当前批次成功创建并显示快照。',
80:'快照创建后，可以继续查看和使用其数据。',
81:'归档成功，所有写入入口均受只读保护。',
82:'取消归档恢复编辑，批次模式标签保持一致。',
83:'筛选正确，实际值、标准及达标判定相互一致。',
84:'学期数据正确筛选，比较标准及判定口径明确。',
85:'预估版正确加载并保留当前教师筛选。',
86:'仅返回指定工号的教师，汇总值正确。',
87:'课程明细与教师工作量汇总一致。',
88:'仅显示符合指定课程代码的明细。',
89:'未填写必填信息时拦截申请提交。',
90:'可以选择模拟附件，并在提交后核验附件保存。',
91:'有效教师申请提交成功，生成可审批记录。',
92:'申请可审核或驳回，通过后正确回写调整结果。',
93:'授权角色可查看并维护对应学年的减免资料。',
94:'减免增改导入删除可完成，达标与结转联动正确。',
95:'学期差额按对应标准计算，说明与来源期间一致。',
96:'全年结转正确入账，后续查询及联动可核验。',
97:'待办刷新正确，无选中项时提交按钮禁用。',
98:'提交计划生成对应通知，内容与提交资料一致。',
99:'通知标记已读后，状态和阅读时间正确保存。',
100:'通知可以跳转到对应业务处理入口。',
101:'首页按学年更新，待办入口的数据范围明确。',
102:'仅清除隔离Term2的本轮复制课程。',
103:'仅清除隔离Term1的本轮导入课程。',
104:'无开课时明确提示，并清除旧模拟结算记录。',
105:'仅清理隔离学期的模拟SIS数据，查询结果为空。',
106:'删除本轮模拟基础课程后，查询结果为空。',
107:'本轮模拟账号及关联教师清理完成。',
108:'仅删除本轮两条模拟课程，并保留审批记录。',
109:'测试草稿学年可删除，或明确提示关联数据限制。',
110:'启用及切换符合学年状态规则，关联配置正确生效。',
}
assert len(expected)==110

def conclusion(s):
    n=s['number'];st=s['state']
    if n==30:return '模板下载需补验；已提供模板继续导入。'
    if n==90:return '附件选择通过；随申请保存待验证。'
    if st=='已完成':return '通过。'
    if st=='已绕行继续':return '采用绕行方式后完成，直接进入的学期加载待完善。'
    if st=='格式调整后继续':return '格式校验通过，调整文件后继续。'
    if '未执行' in st:return '未执行，上游阻塞。' if '阻塞' in st else '未执行，需补充验收。'
    if '阻塞' in st:return '阻塞。'
    if '写权限' in st:return '阻塞，当前账号只读。'
    if '口径' in st or '格式需确认' in st:return '待确认。'
    if '补充文件' in st:return '需补充实际文件验证。'
    return '有待完善项。'

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.25)
def setfont(style,name,size,bold=False):
    st=doc.styles[style];st.font.name=name;st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    rpr=st._element.get_or_add_rPr();rpr.rFonts.set(qn('w:eastAsia'),name)
    for b in st._element.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
setfont('Normal','DengXian',11)
setfont('Title','DengXian Light',25)
setfont('Heading 1','Microsoft YaHei',17,True)
setfont('Heading 2','DengXian',14,True)
setfont('Heading 3','DengXian',12,True)
setfont('Subtitle','DengXian',11)
doc.styles['Subtitle'].font.italic=False
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[name].paragraph_format;pf.line_spacing=1.1;pf.widow_control=True
    pf.space_after=Pt(4)
for nm in ['Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[nm].paragraph_format;pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(5)
for nm in ['TOC 1','TOC 2']:
    if nm not in doc.styles:doc.styles.add_style(nm,1)
    setfont(nm,'DengXian',11,nm=='TOC 1')
    doc.styles[nm].paragraph_format.space_after=Pt(0)
    doc.styles[nm].paragraph_format.line_spacing=1.0
    if nm=='TOC 2':doc.styles[nm].paragraph_format.left_indent=Inches(.22)
for st in doc.styles:
    for border in st._element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
header=sec.header.paragraphs[0];header.text='教师工作量系统   业务流程验收报告'
header.runs[0].font.name='DengXian';header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('626B73')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('626B73')

doc.add_paragraph('教师工作量系统业务流程验收报告','Title')
doc.add_paragraph('UAT环境   2026年10月14日','Subtitle')
p=doc.add_paragraph();p.add_run('验收说明').bold=True
doc.add_paragraph('本科及研究生开课计划均已完成助理提交、主任初审、学院行政复审、行政主管复审和主管领导终审。以下按业务流程列出各步骤的场景、预期、实际结果及结论，保留对应截图。阻塞及未执行环节在该步骤明确标注。')
doc.add_paragraph('本报告复用已完成的操作记录与截图，未重新操作系统。系统管理下的用户、角色、菜单、部门、岗位、字典不纳入验收。')
p=doc.add_paragraph();p.add_run('数据范围').bold=True
doc.add_paragraph('开课审批使用2026—2027学年Summer；资料配置、SIS及结算主要使用2027—2028隔离测试学年；统计查询使用2024—2025资料。各环节分别留证，未将不同批次写作同一批次的连续完成证明。')
doc.add_paragraph('通过Word导航窗格可展开业务流程与步骤。截图保留原始图片数据，可放大查看。')
doc.add_page_break()
p=doc.add_paragraph();r=p.add_run('目录');r.bold=True;r.font.size=Pt(18)
p=doc.add_paragraph()
toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
run=OxmlElement('w:r');txt=OxmlElement('w:t');txt.text='目录在Word中自动更新';run.append(txt);toc.append(run);p._p.append(toc)

def fieldp(label,value,last=False):
    p=doc.add_paragraph();pf=p.paragraph_format
    pf.left_indent=Inches(.8);pf.first_line_indent=Inches(-.8);pf.tab_stops.add_tab_stop(Inches(.8))
    pf.space_after=Pt(2.5);pf.line_spacing=1.04;pf.keep_with_next=True
    p.add_run(label).bold=True;p.add_run('\t'+value)
    return p

mapping=[]
crop_audit={}
def addphoto(s,large=False):
    f=ROOT/records[s['photo']]['screenshot']
    with Image.open(f) as image:
        w,hgt=image.size
        a=np.asarray(image.convert('RGB'),dtype=np.int16)
    # Crop the display frame only. The original JPEG remains unchanged in DOCX.
    # Detect the last row of visible text in the business panel and retain padding.
    d=np.max(np.abs(a[:,1:]-a[:,:-1]),axis=2)
    dark=np.max(a[:,:-1],axis=2)<190
    rowcounts=((d>65)&dark)[:,max(275,int(w*.14)):].sum(axis=1)
    hits=np.flatnonzero(rowcounts>10)
    bottom=min(hgt,max(int(hgt*.55),int(hits[-1])+50)) if len(hits) else hgt
    if large:bottom=hgt
    crop_audit[s['photo']]={'width':w,'height':hgt,'display_bottom':bottom}
    width=min(6.9,(4.5 if large else 3.3)*w/bottom)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9)
    shape=p.add_run().add_picture(str(f),width=Inches(width))
    if bottom<hgt:
        fill=shape._inline.graphic.graphicData.pic.blipFill
        rect=OxmlElement('a:srcRect');rect.set('b',str(round((hgt-bottom)/hgt*100000)))
        fill.insert(1,rect)
        shape.height=Inches(width*bottom/w)
for major,(title,subs) in enumerate(structure,1):
    h=doc.add_heading(f'{major}. {title}',level=1)
    h.paragraph_format.page_break_before=major!=2
    for sub,(subtitle,ids) in enumerate(subs,1):
        if major in [1,2]:
            for i in ids:
                s=steps[i];num=f'{major}.{sub}'
                doc.add_heading(f'{num} {subtitle}',level=2)
                fieldp('验收场景',s['how']);fieldp('预期结果',expected[i]);fieldp('实际结果',s['result']);fieldp('结论',conclusion(s))
                addphoto(s,large=True)
                mapping.append({'number':num,'old_step':i,'title':subtitle,'conclusion':conclusion(s)})
            continue
        doc.add_heading(f'{major}.{sub} {subtitle}',level=2)
        for pos,i in enumerate(ids,1):
            s=steps[i];num=f'{major}.{sub}.{pos}'
            doc.add_heading(f'{num} {s["title"]}',level=3)
            fieldp('验收场景',s['how']);fieldp('预期结果',expected[i]);fieldp('实际结果',s['result']);fieldp('结论',conclusion(s))
            addphoto(s)
            mapping.append({'number':num,'old_step':i,'title':s['title'],'conclusion':conclusion(s)})

out=OUT/'教师工作量系统业务流程验收报告_分级排版版.docx'
doc.save(out)
(OUT/'分级排版映射.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'分级截图显示范围.json').write_text(json.dumps(crop_audit,ensure_ascii=False,indent=2),encoding='utf-8')
print('steps',len(mapping),'main sections',len(structure),'output',out.name)
