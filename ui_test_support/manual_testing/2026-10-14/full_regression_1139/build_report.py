from pathlib import Path
import json, html, collections, csv, zipfile, hashlib
from datetime import datetime
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
records = json.loads((ROOT/'执行结果.json').read_text(encoding='utf-8-sig'))
rows = [r for r in records if not r['id'].startswith('SYS-')]
byid = {r['id']:r for r in records}
counts = collections.Counter(r['status'] for r in rows)
changes = json.loads((ROOT/'数据变更台账.json').read_text(encoding='utf-8-sig'))

# Findings are grouped by observed behavior, without asserting an unverified implementation cause.
specs = [
('D01','P1','特殊审批无法提交','SPECIAL-003 SPECIAL-003R','填写申请人、学部、学年、科研转换、1小时、说明和模拟附件后提交；存在的模拟教师也复测失败。','提示教师用户ID不能为空，而表单无对应选择入口。申请未创建，后续审批受阻。','应能选择有效教师并提交申请，失败提示应能指导用户修正。'),
('D02','P1','运算失败记录仍有正常结果','SET-ENGINE-03','导入零学分且有2小时或1小时的模拟行，执行计算并等待最终完成。','日志称无法运算、成功0条，但行状态NORMAL并有最终工作量2和1，可进入保存与汇总。','无法运算记录应明确失败，不能作为有效工作量参与保存汇总。'),
('D03','P1','新导入记录显示未发生的核实历史','SET-DETAIL-01','打开当天导入的2027年度模拟课程详情。','历史显示2026年6月同步、计算和学部核实通过，本轮未执行核实。','历史应只展示当前记录真实发生的事件与时间。'),
('D04','P1','人工调整覆盖原始课时','SET-ADJUST-02','仅将最终工作量2改4、课程数改0.5，保留原上课时间。','原始课时列也从2变4，与仍为2小时的时间不符。','原始时间解析值应保留，调整结果应独立存储和展示。'),
('D05','P1','重新计算后全局保存丢失说明','SET-COMPARE-02','行级保存偏差说明并刷新确认，再重新计算、全局保存。','重新计算时说明仍在，全局保存后说明清空。','未主动清除的已保存说明应保留。'),
('D06','P1','结算快照创建失败','SET-SNAPSHOT-01','选择已有3条保存记录的隔离批次，输入快照名称并创建。','提示缺少sourceBatchId及创建快照失败。','应带上当前批次信息创建快照，或提供可操作的校验提示。'),
('D07','P1','归档只读仍允许行级写入','SET-ARCHIVE-02','归档模拟批次，在全局只读状态修改行级偏差说明并保存、刷新。','全局保存禁用，但行级说明写入成功并持久化。','归档只读应覆盖全部写入入口。'),
('D08','P1','教师维护找不到刚创建的教师','TEACHER-02 TEACHER-03 TEACHER-IMPORT-02','查询手工创建或导入的模拟教师，再编辑姓名或移除教师身份。','列表存在，提交却返回教师不存在；两种创建方式均复现。','新建和导入教师应能通过教师管理完成后续维护。'),
('D09','P1','结算生成计划丢失计划课时','SET-PLAN-02','从2门已核对计划课时的模拟课程重新生成结算计划。','正确拆为3条，但计划学时和工作量均为--；缺时间行还显示NORMAL无需填写。','应保留开课计划的3、1.5、1.5课时，并正确标识缺失信息。'),
('D10','P2','开课学期下拉依赖先访问配置页','FLOW-PG-00 FLOW-PG-00B CONFIG-TERM-01 FLOW-UG-05','助理重新登录后直接开课，展开具体学期，再对照学期管理。','仅全年，基础学期存在；先访问学期管理后恢复。本科研究生均出现。','开课页面应独立加载已配置学期。当前有绕行方案，审批链已完成。'),
('D11','P2','课程保存后版本记录为空','COURSE-005','新建课程并修改标题保存，再打开版本。','仍提示暂无版本记录、首次保存后会生成。','保存后应产生版本记录，或更正文案和功能约定。'),
('D12','P2','手工课程列表课时未回显','COURSE-009','保存2小时安排，核对编辑、计划和列表。','编辑及计划为2，列表讲座和辅导课时均为-。','列表应展示与安排一致的对应环节课时。'),
('D13','P2','规则列表显示值与条件测试不符','CONFIG-RULE-04','草稿版本添加频率1.25规则并做条件测试。','测试返回1.25，列表却写每周1.0。','列表应反映当前保存值。'),
('D14','P2','规则影响预览范围不符','CONFIG-RULE-06','草稿2027版本新增唯一频率条件，打开影响预览。','预览包含2026模拟课程，课程频率每周且不匹配唯一条件。','预览应按版本与匹配条件限制范围，或明确展示全量参考的含义。'),
('D15','P2','取消归档后模式标签矛盾','SET-ARCHIVE-04','对预估版模拟批次归档再取消归档。','顶部显示期末实际，批次名和页签仍为预估版。','批次模式标签应一致。'),
('D16','P2','跨学年结转把学期标准称为全年','COURSE-CARRY-02','选Term1并查看教师结转详情。','显示合同全年标准6，实际隔离版本全年标准12，6为学期值。','具体学期应标为学期标准，来源期间标签应随筛选改变。'),
('D17','P2','历史导入缺少本次审计记录','COURSE-HISTORY-01','历史导入解析v005模拟文件后确认导入。','弹窗关闭，仍仅旧2024和2023记录，无本次文件、时间或结果。','本次导入应返回明确结果及可追溯记录。实际历史数据落库待核查。'),
('D18','P2','跨学期复制丢失学科部','COURSE-COPY-03','将隔离Term1两条模拟课程复制到Term2并查询。','源为Applied Psychology，目标两条学科部均为-。','复制应保留学科部归属。'),
('D19','P2','合法合教权重保存失败','SET-WEIGHT-02','缺课时、全额工作量0的模拟课，平均分配为50加50并保存。','合计100但提示成功0条失败1条，无原因。','合法比例应保存；若前置数据不满足，应明确提示并阻止误提交。'),
('D20','P2','草稿学年删除失败','CONFIG-YEAR03','删除本轮草稿学年，确认后整页刷新。','草稿仍存在；控制台记录系统异常及删除失败回退。','应删除草稿或清楚说明关联数据等阻止原因。'),
('D21','P2','评分基础展开说明与配置冲突','CONFIG-COPY01','打开2024版本SCH评分基础权重。','PU行和标题为1，展开说明仍将及格不及格写为0.8。','文案应与PU为1一致；本轮不据此推断正式计算错误。'),
('D22','P2','审批跟踪流程图缺少真实终审节点','TRACK-COPY01','完成本科及研究生四级审批后，在审批跟踪筛选已通过并查看流程图。','图中仅主任、行政和领导三节点，把行政主管放在主管领导终审下，未展示实际领导dean终审记录。','流程图应完整展示本轮实际四级审批的角色、处理人和时间。'),
]
findings=[]
for no,priority,title,ids,steps,actual,expected in specs:
    findings.append(dict(id=no,priority=priority,title=title,records=ids.split(),steps=steps,actual=actual,expected=expected))
(ROOT/'缺陷清单.json').write_text(json.dumps(findings,ensure_ascii=False,indent=2),encoding='utf-8')

coverage=[
('角色与审批','8个账号登录；本科和研究生各完成主任至领导四级审批；本科驳回重提；本科研究生待办隔离','两条主链通过，学期下拉有绕行缺陷'),
('开课维护及导入','新增编辑查询复制删除；v005两行导入；每周隔周混合及双教师计划；跨学期复制清除','部分通过，版本、列表、复制及历史导入异常'),
('学年与标准规则','草稿新建、学期编辑、标准复制编辑、规则增改删、匹配与不匹配条件、影响预览','部分通过，未启用草稿；删除失败'),
('课程与教师资料','基础课程增改查删、相同代码导入更新；教师新建与单行导入查询','课程主要操作通过，教师后续维护失败'),
('SIS与结算','Summary重复导入更新、课时格式校验、正常及异常计算、结果保存、人工调整、教师汇总、合教代码维护','部分通过，计算状态、说明及审计等缺陷'),
('归档与快照','批次归档、取消归档、归档行级写保护；创建快照','归档全局控件受控，行级保护失效；快照失败'),
('特殊审批与减免结转','必填、附件选择、实际提交；具体学期结转查询与公式；减免页面权限','特殊审批提交受阻；减免写操作未完成'),
('查询与首页消息','教师工号、明细筛选与合计；实际预估达标页；首页学年切换；本轮消息详情已读跳转','已测操作通过，首页待办口径待确认'),
('导出','教师明细、学部汇总、差异汇总、开课模板下载入口及空特殊审批导出','实际文件未取到，内容与可打开性未验收'),
]
gaps=[
('特殊审批下游','提交提示教师用户ID缺失，没有可用本轮申请。审批通过、驳回和附件持久化均未完成。','修复D01后从创建申请开始走完整审批及调整回写。'),
('快照后续流程','D06导致快照未创建，快照查询、切换、恢复及删除未验证。','修复批次参数后继续快照生命周期。'),
('课时减免写入','管理员和admin_sup均只读，无新增导入编辑入口。未通过改权限绕行。','由产品确认允许写入的角色，再验证增改删、导入及达标结转联动。'),
('导出文件','页面提示导出成功，但浏览器工具没有返回文件；模板由用户直接提供才完成导入。','在普通Edge下载环境复测文件产生、打开、列值、行数、筛选范围。'),
('学年启用与全年结转入账','未启用草稿，避免将原当前学年不可逆归档；仅具体学期查询，未执行全年账本写入。','在可恢复的独立数据环境验证启用、历史状态及全年结转入账。'),
('计算边界及全部组合','本轮覆盖2小时正常、零学分、缺时间、每周隔周及双教师；没有穷举论文指导人数、分钟边界、阈值、各合同及所有规则组合。','补充经业务确认的预期值和边界数据逐项核算。'),
('权限与非功能','验证登录、菜单及审批可见范围；未做直接越权请求、接口安全、压力、跨浏览器或移动端。','单独开展授权的安全及性能兼容测试。'),
]
pending=[r for r in rows if r['status']=='待确认']
manifest=[]
for r in rows:
    p=ROOT/r['screenshot']
    if not p.exists(): raise RuntimeError('缺少截图 '+str(p))
    with Image.open(p) as im: w,h=im.size
    manifest.append(dict(id=r['id'],path=r['screenshot'],width=w,height=h,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
(ROOT/'截图索引.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

def e(s): return html.escape(str(s),quote=True)
def evidence(ids):
    return ' '.join(f'<a href="{e(byid[i]["screenshot"])}" target="_blank">{e(i)}截图</a> <a href="{e(byid[i]["evidence"])}">文字</a>' for i in ids)

parts=['''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>教师工作量系统测试报告 2026年10月14日</title><style>body{font:16px/1.7 "Microsoft YaHei",sans-serif;max-width:1240px;margin:32px auto;padding:0 24px;color:#202124}h1,h2,h3{color:#000;line-height:1.35}table{width:100%;border-collapse:collapse;margin:18px 0}th,td{border:1px solid #d9d9d9;padding:10px;vertical-align:top}th{background:#edf2f7;text-align:left}a{color:#365aa2}details{border:1px solid #ddd;border-radius:4px;margin:10px 0;padding:12px}summary{cursor:pointer;font-weight:600}img{max-width:100%;height:auto}input,select{font:inherit;padding:7px;margin-right:10px}small{color:#555}.fail{color:#b32323}.pass{color:#206538}.note{background:#f4f5f6;padding:12px}nav a{margin-right:20px} @media print{input,select{display:none}details{break-inside:avoid}} </style>''']
parts+=['<h1>教师工作量系统业务测试报告</h1><p>测试日期：2026年10月14日　环境：UAT　方式：Edge中人工逐项操作及截图留证</p><p><b>结论：本轮不满足全部流程通过的验收条件，建议修复关键缺陷并补测后再提交验收。</b>本科和研究生四级开课审批已走通，但特殊审批、结算快照等仍失败。以下只报告实际操作与观察，未执行的流程均单列。</p>', '<nav><a href="#scope">范围</a><a href="#coverage">流程</a><a href="#bugs">缺陷</a><a href="#gaps">未完成项</a><a href="#ledger">数据清理</a><a href="#results">全部记录和原图</a></nav>']
parts+=['<h2 id="scope">范围与统计口径</h2><p>地址：http://hssload-uat.cuhk.edu.cn。账号包含本科助理、研究生助理、学科部主任、本科行政、研究生行政、行政主管、主管领导和管理员。报告不保存任何密码。系统管理的用户、角色、菜单、部门、岗位、字典全部排除，仅允许辅助准备和清理数据；系统配置管理保留在范围内。</p>',f'<p>范围内原始执行记录{len(rows)}条：通过{counts["通过"]}，失败{counts["失败"]}，待确认{counts["待确认"]}，受阻及历史阻塞{counts["受阻"]+counts["阻塞"]}。另有{len(records)-len(rows)}条范围外辅助操作不计入统计。失败记录按行为合并为{len(findings)}项缺陷，建议P1共{sum(f["priority"]=="P1" for f in findings)}项、P2共{sum(f["priority"]=="P2" for f in findings)}项。</p><p>这些是操作记录数量，包含复测、清理与同一问题的多次观察，不能换算为全部功能覆盖率。历史的附件权限阻塞已解除，学期下拉已绕行；原记录保留。旧日期的自动化测试结果不计入本轮，没有运行Selenium或pytest。</p>']
parts+=['<h2 id="coverage">业务流程执行情况</h2><table><tr><th>模块</th><th>实际覆盖</th><th>结论</th></tr>']
for c in coverage:parts.append('<tr>'+''.join('<td>'+e(x)+'</td>' for x in c)+'</tr>')
parts.append('</table><h2 id="bugs">缺陷清单</h2><p>P1为建议优先修复的业务阻断或结果可靠性问题，P2为其他功能及说明问题；最终级别由项目组确认。每项均提供复现条件、实际结果、期望和截图。</p>')
for f in findings:
    parts.append(f'<details><summary>{f["id"]}　{f["priority"]}　{e(f["title"])}</summary><p>复现：{e(f["steps"])}</p><p>实际：{e(f["actual"])}</p><p>期望：{e(f["expected"])}</p><p>{evidence(f["records"])}</p><a href="{e(byid[f["records"][0]]["screenshot"])}" target="_blank"><img loading="lazy" src="{e(byid[f["records"][0]]["screenshot"])}" alt="{e(f["title"])} 原始截图"></a></details>')
parts.append('<h2 id="gaps">未完成流程与待确认事项</h2><table><tr><th>流程</th><th>限制或阻断</th><th>后续验收</th></tr>')
for c in gaps:parts.append('<tr>'+''.join('<td>'+e(x)+'</td>' for x in c)+'</tr>')
parts.append('</table><p>以下待确认记录不作为已确认缺陷，也不算通过。</p>')
for r in pending:parts.append(f'<p><b>{e(r["id"])} {e(r["flow"])}</b>　{e(r["actual"])}　{evidence([r["id"]])}</p>')
parts.append('<h2 id="ledger">模拟数据清理与保留项</h2><p>本轮标记QA1014_1139。2026 Summer两条已审批模拟课程按唯一标题和行标识逐条删除；2027隔离Term1和Term2开课、SIS行、生成结果已清空；QA1014T和QA1014I测试账号已通过用户管理辅助删除，基础课程及自建规则已删除。审批通知和审计记录保留。2027草稿学年删除失败，及其学期、标准、版本和空结算批次可能仍保留，需要开发协助清理。没有启用草稿或切换当前学年；当前2024版本仍24条规则、DPF0.8、PU1、等级制1。本轮没有对全部原规则做逐值快照比对。</p><table><tr><th>数据</th><th>动作</th><th>当前记录</th></tr>')
for c in changes:parts.append('<tr><td>'+e(c.get('key',c.get('marker','')))+'<br>'+e(c.get('year',''))+' '+e(c.get('term',''))+'</td><td>'+e(c.get('action',c.get('type','')) )+'</td><td>'+e(c.get('cleanup',c.get('status','')))+'</td></tr>')
parts.append('</table><h2 id="results">全部范围内执行记录与原始截图</h2><p>截图按原始尺寸保存，点击图片或“原图”可在新标签页放大检查；未人为缩小、模糊或重画截图。每张分辨率及校验值见截图索引.json。</p><input id="q" placeholder="筛选编号或内容"><select id="s"><option value="">全部状态</option>'+''.join(f'<option>{e(k)}</option>' for k in counts)+'</select><span id="n"></span>')
for r in rows:
    mi=next(m for m in manifest if m['id']==r['id'])
    parts.append(f'<details class="record" data-status="{e(r["status"])}"><summary>{e(r["id"])}　{e(r["status"])}　{e(r["module"])}　{e(r["flow"])}</summary><p>{e(r["actual"])}</p><p>{evidence([r["id"]])}　{mi["width"]} × {mi["height"]}</p><a href="{e(r["screenshot"])}" target="_blank"><img loading="lazy" src="{e(r["screenshot"])}" alt="{e(r["id"])}原始截图"></a></details>')
parts.append('''<script>const q=document.querySelector('#q'),s=document.querySelector('#s');function filter(){let n=0;document.querySelectorAll('.record').forEach(x=>{const ok=x.textContent.toLowerCase().includes(q.value.toLowerCase())&&(!s.value||x.dataset.status===s.value);x.hidden=!ok;n+=ok});document.querySelector('#n').textContent=n+' 条记录'}q.oninput=s.onchange=filter;filter()</script></html>''')
(ROOT/'测试报告及截图.html').write_text('\n'.join(parts),encoding='utf-8')
with (ROOT/'范围内执行明细.csv').open('w',encoding='utf-8-sig',newline='') as f:
    wr=csv.DictWriter(f,fieldnames=['id','module','flow','status','actual','evidence','screenshot']);wr.writeheader()
    for r in rows:wr.writerow({k:r.get(k,'') for k in wr.fieldnames})

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[nm];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for st in doc.styles:
    for border in st._element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.16
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
settings=doc.settings.element
noCompress=OxmlElement('w:doNotAutoCompressPictures');settings.append(noCompress)
header=sec.header.paragraphs[0];header.text='教师工作量系统   UAT业务测试报告';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026年10月14日  |  第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9)
def p(t,bold=False):
    z=doc.add_paragraph();z.add_run(t).bold=bold;return z
def heading(t):doc.add_heading(t,level=1)
def table(headers,data,widths,font=10):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for vals in data:
        for i,x in enumerate(vals):t.add_row() if False else None
        cells=t.add_row().cells
        for i,x in enumerate(vals):cells[i].text=str(x)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
    pr.append(b)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        z=OxmlElement('w:'+side);z.set(qn('w:w'),'85');z.set(qn('w:type'),'dxa');margins.append(z)
    pr.append(margins)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF5');cell._tc.get_or_add_tcPr().append(shade)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.05
                for rr in pp.runs:rr.font.size=Pt(font);rr.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def page():doc.add_page_break()

doc.add_paragraph('教师工作量系统业务测试报告','Title')
p('2026年10月14日   UAT环境   Edge人工测试')
heading('测试结论')
p('本轮不满足全部流程通过的验收条件。建议先修复关键缺陷并补齐受阻流程，再提交验收。',True)
p('本科和研究生开课计划均完成四级审批，本科还完成驳回后重建重提。特殊审批无法提交、结算快照创建失败，且发现计算失败状态与结果矛盾、归档后仍可行级写入、说明丢失等问题，不能判定系统无缺陷。')
heading('环境与范围')
p('测试地址 http://hssload-uat.cuhk.edu.cn。测试使用用户提供的8个角色账号，在Edge实际操作并保存截图，没有运行Selenium或pytest。登录重定向后再次输入登录地址进入系统登录页。')
p('纳入：首页、教学计划、结算、待办、系统消息及系统配置管理。按用户最新要求，系统管理下用户、角色、菜单、部门、岗位、字典均不纳入测试，仅辅助准备和清理数据。')
table(['统计项','数量','口径'],[
('范围内执行记录',len(rows),'含复测、清理及历史受阻记录'),('通过 / 失败',f'{counts["通过"]} / {counts["失败"]}','失败记录按行为合并，不等于缺陷数'),('待确认 / 受阻及历史阻塞',f'{counts["待确认"]} / {counts["受阻"]+counts["阻塞"]}','未完成和不确定项目不计为通过'),('合并后缺陷',len(findings),'建议P1 9项，P2 13项'),('范围外辅助记录',len(records)-len(rows),'全部排除出统计和业务缺陷清单')],[2.0,1.0,4.0])
p('统计是本轮操作记录数，不能代表全部功能覆盖率或发布通过率。历史附件权限阻塞已解除，学期下拉已绕行；旧日期自动化结果不计入本轮。')
page();heading('各业务流程执行情况')
table(['模块','实际执行及结果'],[(c[0],c[1]+'。'+c[2]+'。') for c in coverage],[1.2,5.8],10.5)
p('角色链路：助理提交 → 学科部主任 → 对应本科或研究生行政 → 学院行政主管 → 主管领导。两条终审记录均已通过且已处理，最终状态证据见FLOW-FINAL-RECHECK。早期确认弹窗截图不单独作为完成证明。')
page();heading('优先修复的业务缺陷')
p('以下P1为建议优先级，表示业务阻断或结果可靠性受影响。最终级别由项目组确认；完整步骤和原图见配套HTML报告。')
table(['编号','问题及实际表现','证据编号'],[(f['id'],f['title']+'。'+f['actual'],'\n'.join(f['records'])) for f in findings if f['priority']=='P1'],[.5,4.7,1.8],10.5)
page();heading('其他功能及说明缺陷')
table(['编号','问题及实际表现','证据编号'],[(f['id'],f['title']+'。'+f['actual'],'\n'.join(f['records'])) for f in findings if f['priority']=='P2'],[.5,4.7,1.8],10)
page();heading('未完成项目及补测条件')
table(['流程','限制与后续验证'],[(c[0],c[1]+' '+c[2]) for c in gaps],[1.2,5.8],10.5)
heading('需产品确认的口径')
p('首页待办2项与主管待办0条的数据范围是否不同；审批通过后主管是否仍允许删除课程；英文Yes是否为允许的导入值；学期达标是否应比较全年标准。上述事项未直接判为已确认缺陷。')
page();heading('数据清理与验收建议')
p('本轮统一使用QA1014_1139标记。2026 Summer两条已审批模拟课程已按唯一标题及行标识逐条删除，审批和通知审计记录保留。2027隔离Term1、Term2开课、SIS模拟行和生成结果已清空；模拟基础课程、规则、QA1014T及QA1014I测试账号已清理。')
p('残留：QA1014_1139测试学年仍为草稿，删除失败，见D20。其复制学期、标准、版本及空结算批次需开发核查和清理；没有启用草稿，也没有切换原当前学年。')
p('原业务2024版本仍为24条规则，DPF为0.8、PU和等级制为1。本轮规则及标准的写操作在隔离草稿中进行；没有对原版本全部规则做逐值快照比对，不能据此宣称已完成所有规则保真验证。')
p('建议验收前修复P1并复测；补齐特殊审批、快照生命周期、减免写入和导出文件验证；确认待办及锁定权限口径。学年启用与全年结转应在可恢复的独立数据环境验证。')
heading('报告与证据使用方式')
p('Word用于提交和阅读结论；配套“测试报告及截图.html”提供22项缺陷完整步骤、全部范围内执行记录及可放大的原始截图。请解压完整证据包后打开HTML，保留screenshots和evidence目录的相对位置。')
p('截图索引.json记录各图原始分辨率及SHA256；范围内执行明细.csv可直接用Excel筛选；数据变更台账.json记录创建和清理情况。密码未写入报告。')
heading('后续关闭条件')
p('关键缺陷有修复版本、相同模拟条件复测通过，且未完成流程取得实测证据后，更新报告结论。当前结论为业务回归部分通过，不能作为全部功能跑通的验收证明。',True)
page();heading('已通过流程截图证据')
for id_,label in [('FLOW-FINAL-RECHECK','本科与研究生审批最终状态'),('COURSE-PLAN-02','混合频率与双教师分摊计划核对')]:
    pp=p(label+'  '+id_,True);pp.paragraph_format.keep_with_next=True
    doc.add_picture(str(ROOT/byid[id_]['screenshot']),width=Inches(7.0))
    pp=p('原始截图：'+byid[id_]['screenshot']+'。可在配套HTML点击原图放大。');pp.runs[0].font.size=Pt(9)
out=ROOT/'教师工作量系统业务测试报告_20261014.docx';doc.save(out)

print(json.dumps(dict(records=len(records),in_scope=len(rows),counts=dict(counts),defects=len(findings),screenshot_sizes=dict(collections.Counter(f'{m["width"]}x{m["height"]}' for m in manifest)),docx=str(out)),ensure_ascii=False))
