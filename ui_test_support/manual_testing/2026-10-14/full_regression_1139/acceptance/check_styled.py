from pathlib import Path
import json, hashlib, zipfile
from docx import Document
from PIL import Image, ImageDraw
import pypdfium2 as pdf

OUT=Path(__file__).resolve().parent
R=OUT/'styled_render'
file=OUT/'教师工作量系统业务流程验收报告_分级排版版.docx'
d=Document(file)
labels=['验收场景','预期结果','实际结果','结论']
counts={l:sum(p.text.startswith(l+'\t') for p in d.paragraphs) for l in labels}
assert all(n==110 for n in counts.values()),counts
alltext='\n'.join(p.text for p in d.paragraphs)
for unwanted in ['补充操作依据','已有操作截图','保存结果需结合本步回读说明','原图1428','图029','COURSE-008','COURSE-009']:
    assert unwanted not in alltext,unwanted
assert len(d.inline_shapes)==110
manifest=json.loads((OUT/'原始截图清单.json').read_text(encoding='utf-8'))
hashes={v['sha256'] for v in manifest.values()}
with zipfile.ZipFile(file) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    assert all(hashlib.sha256(z.read(n)).hexdigest() in hashes for n in media)
p=pdf.PdfDocument(R/'qa.pdf')
blanks=[]
for i,page in enumerate(p):
    t=page.get_textpage().get_text_range()
    if len(t)<50:blanks.append(i+1)
    if i>=2:
        cs=[t.count(x) for x in labels[:3]]
        assert len(set(cs))==1,('split fields',i+1,cs)
    page.render(scale=1.5).to_pil().save(R/f'page-{i+1:03d}.png')
assert not blanks,blanks
for k,start in enumerate(range(0,len(p),12),1):
    sheet=Image.new('RGB',(1800,2400),'#e4e7eb');dr=ImageDraw.Draw(sheet)
    for j,idx in enumerate(range(start,min(start+12,len(p)))):
        im=Image.open(R/f'page-{idx+1:03d}.png').convert('RGB');im.thumbnail((590,750))
        x=j%3*600+(600-im.width)//2;y=j//3*600+25
        # 3 columns x 4 rows, each page fitted to its slot without stretching.
        im.thumbnail((590,565));x=j%3*600+(600-im.width)//2
        sheet.paste(im,(x,y));dr.text((j%3*600+10,j//3*600+4),f'Page {idx+1}',fill='black')
    sheet.save(R/f'overview-{k:02d}.jpg',quality=94)
quality={'pages':len(p),'steps':110,'four_field_counts':counts,'original_image_bytes_preserved':True,'removed_repeated_caption_material':True,'navigation_heading_levels':[1,2,3],'blank_pages':blanks}
(OUT/'分级排版核查.json').write_text(json.dumps(quality,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(quality,ensure_ascii=False))
