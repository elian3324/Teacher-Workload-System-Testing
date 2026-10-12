from ui_test_support.bootstrap import use_project_python
use_project_python()
import pytest
from openpyxl import load_workbook
from ui_test_support.settings import REPORTS


@pytest.mark.export
@pytest.mark.parametrize("page_id,header", [("teacher", "工号"), ("special_course", "课程编号"),
    ("config_standard", "合同类型")])
def test_export_downloaded_workbook(page, page_id, header):
    page.open(page_id)
    folder = REPORTS / "downloads"
    folder.mkdir(parents=True, exist_ok=True)
    before = {p.name: (p.stat().st_size, p.stat().st_mtime_ns) for p in folder.iterdir() if p.is_file()}
    page.click("导出")
    def completed(d):
        if any(folder.glob("*.crdownload")):
            return False
        for path in folder.glob("*.xlsx"):
            if path.stat().st_size and before.get(path.name) != (path.stat().st_size, path.stat().st_mtime_ns):
                return path
        return False
    downloaded = page.wait.until(completed, "导出没有生成完整的 Excel 文件")
    book = load_workbook(downloaded, read_only=True, data_only=True)
    try:
        values = [str(v) for row in book.active.iter_rows(min_row=1, max_row=5, values_only=True)
                  for v in row if v is not None]
        assert any(header in value for value in values), f"导出文件缺少字段：{header}"
        assert book.active.max_row >= 2
    finally:
        book.close()


if __name__ == "__main__":
    import sys
    sys.dont_write_bytecode = True
    from run_tests import main
    raise SystemExit(main([__file__, *sys.argv[1:]]))
