from pathlib import Path
from collections import OrderedDict
import json, html, base64, csv, hashlib
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent
data=json.loads((OUT/'业务流程信息梳理.json').read_text(encoding='utf-8'))
steps=data['steps']
groups=OrderedDict()
for s in steps:groups.setdefault(s['section'],[]).append(s)
records={r['id']:r for r in json.loads((ROOT/'执行结果.json').read_text(encoding='utf-8-sig')) if not r['id'].startswith('SYS-')}
images={}
for s in steps:
    for id_ in set(s['records']+[s['photo']]):
        if id_ in images:continue
        f=ROOT/records[id_]['screenshot']
        with Image.open(f) as im:w,h=im.size
        images[id_]=dict(width=w,height=h,data='data:image/jpeg;base64,'+base64.b64encode(f.read_bytes()).decode(),sha256=hashlib.sha256(f.read_bytes()).hexdigest())

def e(x):return html.escape(str(x),quote=True)
def img(id_,caption):
    im=images[id_]
    return f'<figure><img loading="lazy" src="{im["data"]}" alt="{e(caption)}" tabindex="0" onclick="showImage(this)" onkeydown="if(event.key===\'Enter\')showImage(this)"><figcaption>{e(caption)} · 原图 {im["width"]} × {im["height"]} · 点击放大</figcaption></figure>'

intro='开课计划由助理提交，经学科部主任、对应学院行政、学院行政主管、主管领导四级审批，最终通过的主流程已经完成。本报告按业务顺序逐步展示已有操作和截图；其他环节在对应步骤标明完成情况及阻塞位置。'
context='开课审批使用2026至2027学年Summer；资料配置、SIS及结算主要使用2027至2028隔离草稿；工作量和达标查询使用2024至2025资料。各段使用不同模拟数据，现有证据不等同于同一批次从开课到结算归档连续贯通。'
scope='本次依据2026年10月14日已经完成的Edge操作重新整理，没有再次操作浏览器。系统管理下的用户、角色、菜单、部门、岗位和字典不纳入验收，只在准备或清理数据时辅助使用。'
parts=['''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>教师工作量系统业务流程验收报告</title><style>
:root{color-scheme:light}*{box-sizing:border-box}body{margin:0;font:16px/1.75 "Microsoft YaHei","PingFang SC",sans-serif;color:#242a30;background:#f4f6f8}header{background:white;border-bottom:1px solid #d9dfe5;padding:30px max(28px,calc((100vw - 1380px)/2));}h1,h2,h3{color:#000;line-height:1.45}h1{font-size:30px;margin:0 0 10px}header p{max-width:1200px;margin:8px 0}.meta,.caption{color:#5e6670;font-size:14px}.layout{max-width:1440px;display:grid;grid-template-columns:250px minmax(0,1fr);gap:24px;margin:24px auto;padding:0 24px}aside{position:sticky;top:15px;align-self:start;max-height:95vh;overflow:auto;background:white;border:1px solid #d9dfe5;padding:18px}aside a{display:block;color:#34495e;text-decoration:none;margin:8px 0;font-size:14px}aside a:hover{text-decoration:underline}main{min-width:0}section{scroll-margin-top:15px}h2{font-size:23px;margin:6px 0 16px}article{background:white;padding:24px;margin:0 0 24px;border:1px solid #d9dfe5;border-radius:4px;scroll-margin-top:18px}h3{font-size:20px;margin:0 0 10px}.status{display:inline-block;font-size:13px;padding:3px 10px;border:1px solid #c4d7ca;background:#eff7f1;color:#245c39;margin-bottom:12px}.pending{background:#fff7e9;border-color:#e4d0a9;color:#775316}.steptext{margin:6px 0}.steptext strong{font-weight:600}figure{margin:18px 0 6px}img{display:block;width:100%;height:auto;border:1px solid #e0e3e6;cursor:zoom-in}figcaption{color:#626b75;font-size:13px;margin-top:6px}details{border-top:1px solid #e0e3e6;padding-top:12px;margin-top:20px}summary{cursor:pointer;font-size:14px;color:#40566d}.source{margin-top:20px}.source h4{font-size:15px;margin:0}.source p{font-size:14px}.source img{max-height:660px;object-fit:contain;background:#f9fafb}.toolbar{display:flex;gap:10px;margin-bottom:20px}input,select,button{font:inherit;padding:8px 10px;border:1px solid #c4cbd3;border-radius:3px;background:white}input{width:100%}button{cursor:pointer}.hidden{display:none!important}dialog{padding:0;border:0;max-width:98vw;max-height:97vh;width:96vw}dialog::backdrop{background:rgba(0,0,0,.75)}.viewerbar{padding:10px;background:#f4f6f8;display:flex;gap:10px;position:sticky;top:0;z-index:1}.imageviewport{overflow:auto;max-height:85vh;background:#fff}.imageviewport img{width:100%;max-width:100%;border:0;cursor:default}.imageviewport.original img{width:auto;max-width:none}footer{color:#666;padding:30px;text-align:center;font-size:13px}@media(max-width:1000px){.layout{display:block}aside{position:static;max-height:340px;margin-bottom:20px}article{padding:18px}}@media print{body{background:white}.layout{display:block;padding:0;margin:0}aside,.toolbar,.viewerbar{display:none}article{break-inside:avoid;border:0;border-bottom:1px solid #ddd}details{display:none}header{padding:0}img{max-height:220mm;object-fit:contain}}
</style></head><body>''']
parts.append('<header><h1>教师工作量系统业务流程验收报告</h1><p class="meta">验收留证日期 2026年10月14日　环境 UAT　报告形式 逐步说明与原始截图</p><p>'+intro+'</p><p>'+scope+'</p><details><summary>测试数据和证据使用说明</summary><p>'+context+'</p><p>每一步展示简短操作说明、已有结果和对应截图。截图保持原始分辨率，点击后可切换原始尺寸。预览、确认框仅证明当时操作阶段；上游阻塞图不表示下游已经执行。全文不保存密码。</p></details></header><div class="layout"><aside><strong>按业务流程阅读</strong>')
for j,(g,ss) in enumerate(groups.items(),1):parts.append(f'<a href="#section-{j}">{j} {e(g)}<br><small>步骤 {ss[0]["number"]} 至 {ss[-1]["number"]}</small></a>')
parts.append('</aside><main><div class="toolbar"><input id="query" placeholder="查找步骤名称或业务内容"><select id="state"><option value="all">全部步骤</option><option value="open">只看阻塞和需补验步骤</option></select></div>')
for j,(g,ss) in enumerate(groups.items(),1):
    parts.append(f'<section id="section-{j}"><h2>{j} {e(g)}</h2>')
    for s in ss:
        pending= any(k in s['state'] for k in ['阻塞','未执行','补充','补验'])
        ok=s['state']=='已完成'
        parts.append(f'<article id="step-{s["number"]}" data-open="{int(pending)}"><h3>步骤 {s["number"]:03d} {e(s["title"])}</h3><span class="status {"" if ok else "pending"}">{e(s["state"])}</span><p class="steptext"><strong>怎么做</strong>　{e(s["how"])}</p><p class="steptext"><strong>验收记录</strong>　{e(s["result"])}</p>')
        cap=f'图{s["number"]:03d} '+s['title']
        if not s['records']:cap+=' 上游或前置页面留图 本步骤未执行'
        if s['photo']=='FLOW-FINAL-RECHECK':cap+=' 最终结果两条均已通过'
        parts.append(img(s['photo'],cap))
        parts.append('<p class="caption">'+e(s['photo_note'])+'</p>')
        if s['sources']:
            parts.append('<details><summary>展开本步骤的补充截图与操作依据</summary>')
            for r in s['sources']:
                parts.append(f'<div class="source"><h4>{e(r["flow"])}</h4><p>{e(r["actual"])}</p>')
                if r['id']!=s['photo']:parts.append(img(r['id'],r['flow']+' 操作留图'))
                parts.append('<p class="caption">来源编号 '+e(r['id'])+'</p></div>')
            parts.append('</details>')
        parts.append('</article>')
    parts.append('</section>')
parts.append('''</main></div><dialog id="viewer"><div class="viewerbar"><button onclick="document.querySelector('#viewer').close()">关闭</button><button onclick="document.querySelector('.imageviewport').classList.toggle('original')">切换原始尺寸</button><span>可滚动查看原图，按 Esc 关闭</span></div><div class="imageviewport"><img id="large" alt="放大的原始截图"></div></dialog><footer>教师工作量系统业务流程验收报告　依据2026年10月14日已留存证据编写</footer><script>function showImage(i){document.querySelector('#large').src=i.src;document.querySelector('.imageviewport').classList.remove('original');document.querySelector('#viewer').showModal()}const q=document.querySelector('#query'),s=document.querySelector('#state');function filter(){document.querySelectorAll('article').forEach(x=>{x.classList.toggle('hidden',!x.textContent.toLowerCase().includes(q.value.toLowerCase())||(s.value==='open'&&x.dataset.open!=='1'))});document.querySelectorAll('main section').forEach(x=>x.classList.toggle('hidden',!x.querySelector('article:not(.hidden)')))}q.oninput=s.onchange=filter;</script></body></html>''')
htmlfile=OUT/'教师工作量系统业务流程验收报告_逐步原图版.html'
htmlfile.write_text('\n'.join(parts),encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[nm];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for st in doc.styles:
    for b in st._element.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(7)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(12)
no=OxmlElement('w:doNotAutoCompressPictures');doc.settings.element.append(no)
header=sec.header.paragraphs[0];header.text='教师工作量系统  业务流程验收报告'
header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026年10月14日  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(9)
def p(t,bold=False,size=None):
    a=doc.add_paragraph();r=a.add_run(t);r.bold=bold
    if size:r.font.size=Pt(size)
    return a
def h(t):return doc.add_heading(t,level=1)
doc.add_paragraph('教师工作量系统业务流程验收报告','Title')
p('验收留证日期 2026年10月14日   UAT环境   逐步截图版')
h('业务流程验收说明');p(intro)
h('验收范围');p(scope)
h('业务主线');p('准备学年标准和课程资料 → 助理维护并提交开课计划 → 主任及学院各级审批 → 导入SIS与上课时间 → 生成并保存结算结果 → 调整及查询 → 报表快照与归档。')
h('阅读方式');p('正文按业务流程逐步展示，每一步配有操作说明、已有结果和截图。已完成处按事实记录；遇到阻塞就标在该步骤。没有执行的后续操作明确写为未执行，配图仅用于定位上游阻塞。')
p('原图版HTML包含全部关联截图，点击图片可按原始尺寸查看。Word保留原始图片数据，可放大阅读；原图版不依赖其他文件，可直接在Edge中打开。')
h('数据场景');p(context)
doc.add_page_break();h('业务流程目录')
for j,(g,ss) in enumerate(groups.items(),1):
    p(f'{j}  {g}   步骤 {ss[0]["number"]:03d} 至 {ss[-1]["number"]:03d}',True,size=11)
p('本报告依据已有操作证据整理，未使用签字或结论代替未执行步骤。账号口令不写入正文。',size=10)

for s in steps:
    doc.add_page_break()
    p(s['section'],size=10)
    h(f'步骤 {s["number"]:03d} {s["title"]}')
    p('状态  '+s['state'],True)
    p('操作说明  '+s['how'])
    p('验收记录  '+s['result'])
    im=images[s['photo']]
    width=7.2;height=width*im['height']/im['width']
    if height>5.35:width=5.35*im['width']/im['height']
    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(ROOT/records[s['photo']]['screenshot']),width=Inches(width))
    note=f'图{s["number"]:03d}  {s["title"]}  原图{im["width"]}×{im["height"]}'
    if not s['records']:note+='  上游或前置页面 本步骤未执行'
    if s['photo']=='FLOW-FINAL-RECHECK':note+='  两条最终状态均已通过'
    p(note,size=9)
    p(s['photo_note'],size=9)
    if s['records']:p('补充操作依据  '+'、'.join(s['records'])+'。全部关联截图见原图版同一步骤。',size=9)
    else:p('此处需补充实测证据，尚无本步骤执行成功截图。',size=9)

docfile=OUT/'教师工作量系统业务流程验收报告_20261014.docx'
doc.save(docfile)
with (OUT/'业务流程步骤索引.csv').open('w',encoding='utf-8-sig',newline='') as f:
    wr=csv.writer(f);wr.writerow(['步骤','业务阶段','流程','操作说明','验收记录','状态','主图','关联证据'])
    for s in steps:wr.writerow([s['number'],s['section'],s['title'],s['how'],s['result'],s['state'],s['photo'],' '.join(s['records'])])
(OUT/'原始截图清单.json').write_text(json.dumps({k:{q:v for q,v in x.items() if q!='data'} for k,x in images.items()},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'steps':len(steps),'sections':len(groups),'covered_records':data['covered_source_records'],'images':len(images),'html_bytes':htmlfile.stat().st_size,'docx_bytes':docfile.stat().st_size,'expected_word_pages':len(steps)+2},ensure_ascii=False))
